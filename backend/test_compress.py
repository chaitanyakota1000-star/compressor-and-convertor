import os
import io
import shutil
import pytest
from pathlib import Path
from fastapi.testclient import TestClient

# Import the FastAPI app
from app.main import app
from app.database import uploaded_files, compressed_files, users_db
from app.config import UPLOAD_DIR, COMPRESSED_DIR

# Create the test client
client = TestClient(app)

def setup_module(module):
    """Ensure clean directories before running tests"""
    for f in UPLOAD_DIR.glob("*"):
        if f.is_file():
            f.unlink()
    for f in COMPRESSED_DIR.glob("*"):
        if f.is_file():
            f.unlink()
    uploaded_files.clear()
    compressed_files.clear()
    users_db.clear()

def test_health_check():
    response = client.get("/")
    assert response.status_code == 200
    assert response.json()["status"] == "healthy"
    assert response.json()["auth"] == "active"

def test_user_signup_and_login():
    signup_payload = {"email": "testuser@example.com", "password": "securepassword123"}
    resp_signup = client.post("/api/signup", json=signup_payload)
    assert resp_signup.status_code == 200
    assert "Registration successful" in resp_signup.json()["message"]
    
    resp_login = client.post("/api/login", json=signup_payload)
    assert resp_login.status_code == 200
    login_data = resp_login.json()
    assert "token" in login_data

def test_protected_compress_without_token():
    compress_resp = client.post("/api/compress", json={
        "file_ids": ["dummy-id"],
        "quality": 80,
        "archive_format": "zip"
    })
    assert compress_resp.status_code == 401

def test_image_format_conversion():
    from PIL import Image
    user_payload = {"email": "convuser@example.com", "password": "password12345"}
    client.post("/api/signup", json=user_payload)
    token = client.post("/api/login", json=user_payload).json()["token"]

    img = Image.new("RGBA", (100, 100), color="blue")
    img_byte_arr = io.BytesIO()
    img.save(img_byte_arr, format="PNG")
    img_byte_arr.seek(0)
    png_content = img_byte_arr.read()
    
    file_data = {"file": ("convert_test.png", io.BytesIO(png_content), "image/png")}
    file_id = client.post("/api/upload", files=file_data).json()["file_id"]
    
    headers = {"Authorization": f"Bearer {token}"}
    compress_resp = client.post("/api/compress", json={
        "file_ids": [file_id],
        "quality": 80,
        "archive_format": "zip",
        "target_format": "WEBP"
    }, headers=headers)
    
    assert compress_resp.status_code == 200
    comp_data = compress_resp.json()
    assert comp_data["filename"] == "compressed_convert_test.webp"

def test_image_target_size_optimization():
    from PIL import Image
    user_payload = {"email": "optuser@example.com", "password": "password12345"}
    client.post("/api/signup", json=user_payload)
    token = client.post("/api/login", json=user_payload).json()["token"]

    img = Image.new("RGB", (300, 300), color="green")
    img_byte_arr = io.BytesIO()
    img.save(img_byte_arr, format="JPEG")
    img_byte_arr.seek(0)
    jpeg_content = img_byte_arr.read()
    
    file_data = {"file": ("opt_test.jpg", io.BytesIO(jpeg_content), "image/jpeg")}
    file_id = client.post("/api/upload", files=file_data).json()["file_id"]
    
    headers = {"Authorization": f"Bearer {token}"}
    compress_resp = client.post("/api/compress", json={
        "file_ids": [file_id],
        "quality": 90,
        "archive_format": "zip",
        "optimize_size": True,
        "target_min_size": 100,
        "target_max_size": 10000
    }, headers=headers)
    
    assert compress_resp.status_code == 200
    comp_data = compress_resp.json()
    assert comp_data["size"] <= 10000

def test_download_range_requests():
    user_payload = {"email": "rangeuser@example.com", "password": "password12345"}
    client.post("/api/signup", json=user_payload)
    token = client.post("/api/login", json=user_payload).json()["token"]
    
    file_content = b"abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ12345"
    file_data = {"file": ("range_test.txt", io.BytesIO(file_content), "text/plain")}
    file_id = client.post("/api/upload", files=file_data).json()["file_id"]
    
    headers = {"Authorization": f"Bearer {token}"}
    comp_resp = client.post("/api/compress", json={
        "file_ids": [file_id],
        "quality": 80,
        "archive_format": "zip"
    }, headers=headers)
    
    download_id = comp_resp.json()["download_id"]
    
    dl_full_resp = client.get(f"/api/download/{download_id}")
    assert dl_full_resp.status_code == 200
    assert dl_full_resp.headers["Accept-Ranges"] == "bytes"
    total_size = len(dl_full_resp.content)
    
    dl_range_resp = client.get(
        f"/api/download/{download_id}", 
        headers={"Range": "bytes=0-9"}
    )
    assert dl_range_resp.status_code == 206
    assert dl_range_resp.headers["Content-Range"] == f"bytes 0-9/{total_size}"
    assert len(dl_range_resp.content) == 10

# ==========================================
# FILE FORMAT CONVERSION ROUTE TESTS
# ==========================================

def test_convert_images_to_pdf():
    from PIL import Image
    user_payload = {"email": "imgpdfuser@example.com", "password": "password12345"}
    client.post("/api/signup", json=user_payload)
    token = client.post("/api/login", json=user_payload).json()["token"]

    # 1. Upload two dummy images
    img1 = Image.new("RGB", (50, 50), color="red")
    img2 = Image.new("RGB", (50, 50), color="green")
    
    io1 = io.BytesIO()
    io2 = io.BytesIO()
    img1.save(io1, format="JPEG")
    img2.save(io2, format="JPEG")
    io1.seek(0)
    io2.seek(0)
    
    u1 = client.post("/api/upload", files={"file": ("img1.jpg", io1, "image/jpeg")}).json()["file_id"]
    u2 = client.post("/api/upload", files={"file": ("img2.jpg", io2, "image/jpeg")}).json()["file_id"]
    
    # 2. Trigger conversion to PDF
    headers = {"Authorization": f"Bearer {token}"}
    convert_resp = client.post("/api/convert", json={
        "file_ids": [u1, u2],
        "conversion_type": "image_to_pdf"
    }, headers=headers)
    
    assert convert_resp.status_code == 200
    conv_data = convert_resp.json()
    assert conv_data["filename"] == "converted_images.pdf"
    
    # 3. Download and verify is PDF
    dl_resp = client.get(f"/api/download/{conv_data['download_id']}")
    assert dl_resp.status_code == 200
    assert dl_resp.content.startswith(b"%PDF-")

def test_convert_pdf_to_images():
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet
    
    user_payload = {"email": "pdfimguser@example.com", "password": "password12345"}
    client.post("/api/signup", json=user_payload)
    token = client.post("/api/login", json=user_payload).json()["token"]

    # 1. Create a 2-page PDF document in memory using ReportLab
    pdf_buffer = io.BytesIO()
    styles = getSampleStyleSheet()
    pdf_doc = SimpleDocTemplate(pdf_buffer, pagesize=letter)
    story = [
        Paragraph("Page 1 content here.", styles['Normal']),
        PageBreak(),
        Paragraph("Page 2 content here.", styles['Normal'])
    ]
    pdf_doc.build(story)
    pdf_buffer.seek(0)
    pdf_bytes = pdf_buffer.read()
    
    # 2. Upload PDF
    file_id = client.post("/api/upload", files={"file": ("doc2pages.pdf", io.BytesIO(pdf_bytes), "application/pdf")}).json()["file_id"]
    
    # 3. Convert PDF to images
    headers = {"Authorization": f"Bearer {token}"}
    convert_resp = client.post("/api/convert", json={
        "file_ids": [file_id],
        "conversion_type": "pdf_to_image",
        "target_img_format": "PNG"
    }, headers=headers)
    
    assert convert_resp.status_code == 200
    conv_data = convert_resp.json()
    # Since it is a 2-page document, it should return a zip file
    assert conv_data["filename"] == "doc2pages_pages.zip"

def test_convert_pdf_to_docx():
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet
    
    user_payload = {"email": "pdfdocxuser@example.com", "password": "password12345"}
    client.post("/api/signup", json=user_payload)
    token = client.post("/api/login", json=user_payload).json()["token"]

    # 1. Create simple PDF
    pdf_buffer = io.BytesIO()
    styles = getSampleStyleSheet()
    pdf_doc = SimpleDocTemplate(pdf_buffer, pagesize=letter)
    story = [Paragraph("This is the source PDF text to be converted to DOCX.", styles['Normal'])]
    pdf_doc.build(story)
    pdf_buffer.seek(0)
    pdf_bytes = pdf_buffer.read()
    
    # 2. Upload PDF
    file_id = client.post("/api/upload", files={"file": ("input.pdf", io.BytesIO(pdf_bytes), "application/pdf")}).json()["file_id"]
    
    # 3. Convert to DOCX
    headers = {"Authorization": f"Bearer {token}"}
    convert_resp = client.post("/api/convert", json={
        "file_ids": [file_id],
        "conversion_type": "pdf_to_docx"
    }, headers=headers)
    
    assert convert_resp.status_code == 200
    conv_data = convert_resp.json()
    assert conv_data["filename"] == "input.docx"

def test_convert_docx_to_pdf():
    from docx import Document
    user_payload = {"email": "docxpdfuser@example.com", "password": "password12345"}
    client.post("/api/signup", json=user_payload)
    token = client.post("/api/login", json=user_payload).json()["token"]

    # 1. Create a simple .docx in memory using python-docx
    docx_doc = Document()
    docx_doc.add_heading("Test Heading", level=1)
    docx_doc.add_paragraph("This is test paragraph content inside the docx file.")
    
    docx_buffer = io.BytesIO()
    docx_doc.save(docx_buffer)
    docx_buffer.seek(0)
    docx_bytes = docx_buffer.read()
    
    # 2. Upload DOCX
    file_id = client.post("/api/upload", files={"file": ("wordfile.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}).json()["file_id"]
    
    # 3. Convert DOCX to PDF
    headers = {"Authorization": f"Bearer {token}"}
    convert_resp = client.post("/api/convert", json={
        "file_ids": [file_id],
        "conversion_type": "docx_to_pdf"
    }, headers=headers)
    
    assert convert_resp.status_code == 200
    conv_data = convert_resp.json()
    assert conv_data["filename"] == "wordfile.pdf"
    
    # 4. Download and verify it starts with %PDF-
    dl_resp = client.get(f"/api/download/{conv_data['download_id']}")
    assert dl_resp.status_code == 200
    assert dl_resp.content.startswith(b"%PDF-")

if __name__ == "__main__":
    import sys
    sys.exit(pytest.main([__file__, "-v"]))
